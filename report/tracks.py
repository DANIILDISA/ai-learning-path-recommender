"""Track report: JSON and Excel.

Covers the ТЗ requirement "форма отчёта: JSON и Excel". This module knows
about tracks, the catalogue and the data-quality report; it has no idea
model groups exist -- that comparison lives in `benchmark/`.
"""
from __future__ import annotations

import json
import statistics
from pathlib import Path

from openpyxl import Workbook

from catalog.catalog import Catalog
from core.models import Track
from quality.report import DataQualityReport
from report.excel import HEADER_FONT, autosize_columns, write_header_row

SHEET_TRACKS = "Треки"
SHEET_SUMMARY = "Сводка"
SHEET_QUALITY = "Качество данных"

_TRACKS_HEADERS = [
    "Пользователь",
    "Должность",
    "ИОГВ",
    "Грейд",
    "№",
    "Курс",
    "Тип",
    "Часы",
    "Форма",
    "Балл",
    "Обоснования",
]


def track_report_payload(tracks: dict[str, Track]) -> list[dict]:
    """JSON-ready payload: one entry per profile, sorted by `user_id` so the
    same corpus always produces the same file byte-for-byte."""
    return [tracks[user_id].model_dump() for user_id in sorted(tracks)]


def write_json_report(tracks: dict[str, Track], out_path: str | Path) -> None:
    """Write the track report as JSON. Cyrillic is written literally
    (`ensure_ascii=False`), not as `\\uXXXX` escapes."""
    payload = track_report_payload(tracks)
    Path(out_path).write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def _delivery_label(catalog: Catalog, course_key: str) -> str:
    course = catalog.get(course_key)
    if course is None or not course.deliveries:
        return ""
    return "/".join(sorted(d.value for d in course.deliveries))


def _write_tracks_sheet(wb: Workbook, tracks: dict[str, Track], catalog: Catalog) -> None:
    ws = wb.create_sheet(SHEET_TRACKS)
    write_header_row(ws, _TRACKS_HEADERS)
    row = 2
    for user_id in sorted(tracks):
        track = tracks[user_id]
        for number, rec in enumerate(track.recommendations, start=1):
            course = catalog.get(rec.course_key)
            ws.cell(row=row, column=1, value=track.user_id)
            ws.cell(row=row, column=2, value=track.position)
            ws.cell(row=row, column=3, value=track.iogv)
            ws.cell(row=row, column=4, value=track.grade)
            ws.cell(row=row, column=5, value=number)
            ws.cell(row=row, column=6, value=rec.title)
            ws.cell(row=row, column=7, value=_delivery_label(catalog, rec.course_key))
            ws.cell(row=row, column=8, value=course.hours if course else None)
            ws.cell(row=row, column=9, value=(course.study_form if course else None) or "")
            ws.cell(row=row, column=10, value=rec.score)
            ws.cell(row=row, column=11, value="; ".join(rec.reasons))
            row += 1
    autosize_columns(ws)


def _write_summary_sheet(wb: Workbook, tracks: dict[str, Track], catalog: Catalog) -> None:
    ws = wb.create_sheet(SHEET_SUMMARY)
    write_header_row(ws, ["Показатель", "Значение"])

    lengths = [len(track.recommendations) for track in tracks.values()]
    mean_length = statistics.mean(lengths) if lengths else 0

    grade_counts: dict[str, int] = {}
    for track in tracks.values():
        grade_counts[track.grade] = grade_counts.get(track.grade, 0) + 1

    competency_counts: dict[str, int] = {}
    for track in tracks.values():
        for rec in track.recommendations:
            course = catalog.get(rec.course_key)
            label = course.competency_group.value if course and course.competency_group else "Без блока"
            competency_counts[label] = competency_counts.get(label, 0) + 1

    row = 2
    for label, value in [
        ("Профилей", len(tracks)),
        ("Курсов в каталоге", len(catalog)),
        ("Треков построено", len(tracks)),
        ("Среднее число рекомендаций в треке", round(mean_length, 2)),
    ]:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Распределение по грейдам").font = HEADER_FONT
    row += 1
    for grade in sorted(grade_counts):
        ws.cell(row=row, column=1, value=grade)
        ws.cell(row=row, column=2, value=grade_counts[grade])
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Распределение по компетенционным блокам").font = HEADER_FONT
    row += 1
    for label in sorted(competency_counts):
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=competency_counts[label])
        row += 1

    autosize_columns(ws)


def _write_quality_sheet(wb: Workbook, quality: DataQualityReport) -> None:
    ws = wb.create_sheet(SHEET_QUALITY)
    write_header_row(ws, ["Показатель", "Значение"])

    rows = [
        ("Строк истории обучения", quality.history_rows),
        ("Профилей пользователей", quality.profiles),
        ("Курсов в каталоге", quality.catalogue_courses),
        ("Курсов истории без каталога", len(quality.unmatched_keys)),
        ("Курсов в двух форматах", len(quality.dual_delivery_keys)),
        ("Дублирующихся «Результатов»", len(quality.duplicated_outcomes)),
        ("Точных дублей строк истории", quality.duplicate_rows),
        ("Противоречивых статусов", quality.contradictory_statuses),
        ("Написаний должностей (сырых)", quality.raw_positions),
        ("Написаний должностей (после нормализации)", quality.normalised_positions),
    ]
    for offset, (label, value) in enumerate(rows):
        ws.cell(row=2 + offset, column=1, value=label)
        ws.cell(row=2 + offset, column=2, value=value)

    autosize_columns(ws)


def write_excel_report(
    tracks: dict[str, Track],
    catalog: Catalog,
    quality: DataQualityReport,
    out_path: str | Path,
) -> None:
    """Write the three-sheet track report: Треки, Сводка, Качество данных."""
    wb = Workbook()
    # `Workbook()` starts with one default sheet ("Sheet"); drop it once the
    # real sheets exist so the workbook has exactly the three named above.
    default_sheet = wb.active
    _write_tracks_sheet(wb, tracks, catalog)
    _write_summary_sheet(wb, tracks, catalog)
    _write_quality_sheet(wb, quality)
    wb.remove(default_sheet)
    wb.save(out_path)
