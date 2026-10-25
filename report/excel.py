"""Shared openpyxl formatting helpers.

Kept separate from `report/tracks.py` (and, later, `benchmark/report.py`) so
the look of every workbook this product writes -- bold headers, sized
columns, a frozen header row -- comes from one place instead of being
re-implemented per sheet.
"""
from __future__ import annotations

from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FONT = Font(bold=True)


def write_header_row(ws: Worksheet, headers: list[str], row: int = 1) -> None:
    """Write a bold header row and freeze the sheet just below it."""
    for column, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=column, value=text)
        cell.font = HEADER_FONT
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate


def autosize_columns(ws: Worksheet, min_width: int = 8, max_width: int = 60) -> None:
    """Size each column to its widest cell's text, clamped to a sane range.

    A methodologist opens this workbook directly; unsized columns showing
    "###" or truncated Russian labels is not acceptable.
    """
    widths: dict[str, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            letter = cell.column_letter
            widths[letter] = max(widths.get(letter, 0), length)
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = max(min_width, min(max_width, width + 2))
