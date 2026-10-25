"""Comparative Excel workbook across model groups.

This is the ТЗ deliverable "Excel файл с результатами анализа работы группы
ИИ-агентов" and the evidence behind "глубокий сравнительный анализ
отечественных, зарубежных и локальных моделей с обоснованием выбора по
метрикам". It has no idea how the numbers were produced -- `benchmark.runner`
already settled the cache-bypass question by the time a `GroupResult` reaches
here -- this module only renders `GroupResult`/`GroupMetrics` into sheets.

The one rule every sheet here follows: a group with `status="not_measured"`
gets a labelled cell everywhere it would otherwise need a number, never a
blank or a 0 -- see `report/tracks.py`'s `report/excel.py` helpers, reused
below so the two workbooks this product writes look like one system.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from benchmark.metrics import GroupMetrics, agreement_between, choose_group, summarize
from benchmark.runner import GroupResult
from core.config import AppConfig
from report.excel import HEADER_FONT, autosize_columns, write_header_row

SHEET_SUMMARY = "Сводка"
SHEET_METRICS = "Метрики"
SHEET_PROFILES = "По профилям"
SHEET_METHODOLOGY = "Методика"

NOT_MEASURED = "не измерено"

_METRICS_HEADERS = [
    "Группа",
    "Модель",
    "Статус",
    "Задержка (медиана, с)",
    "Задержка (среднее, с)",
    "Токены вход",
    "Токены выход",
    "Валидный JSON",
    "В пределах пула",
    "Воспроизводимость",
    "Согласие с другими группами",
]


def _pct(value: float | None) -> str:
    return f"{value * 100:.0f}%" if value is not None else "—"


def _not_measured_cell(metrics: GroupMetrics) -> str:
    return f"{NOT_MEASURED}: {metrics.reason}" if metrics.reason else NOT_MEASURED


def _write_summary_sheet(wb: Workbook, metrics_by_group: dict[str, GroupMetrics]) -> None:
    ws = wb.create_sheet(SHEET_SUMMARY)
    write_header_row(ws, ["Показатель", "Значение"])
    name, verdict = choose_group(metrics_by_group)

    ws.cell(row=2, column=1, value="Вывод")
    ws.cell(row=2, column=2, value=verdict)
    ws.cell(row=3, column=1, value="Выбранная группа")
    ws.cell(row=3, column=2, value=name or NOT_MEASURED)

    row = 5
    ws.cell(row=row, column=1, value="Статус групп").font = HEADER_FONT
    row += 1
    for group_name in sorted(metrics_by_group):
        metrics = metrics_by_group[group_name]
        ws.cell(row=row, column=1, value=group_name)
        ws.cell(
            row=row, column=2,
            value="измерено" if metrics.status == "measured" else _not_measured_cell(metrics),
        )
        row += 1
    autosize_columns(ws)


def _write_metrics_sheet(
    wb: Workbook, metrics_by_group: dict[str, GroupMetrics], results: dict[str, GroupResult],
) -> None:
    ws = wb.create_sheet(SHEET_METRICS)
    write_header_row(ws, _METRICS_HEADERS)

    names = sorted(metrics_by_group)
    row = 2
    for group_name in names:
        metrics = metrics_by_group[group_name]
        ws.cell(row=row, column=1, value=group_name)
        ws.cell(row=row, column=2, value=metrics.model or "—")

        if metrics.status != "measured":
            ws.cell(row=row, column=3, value=_not_measured_cell(metrics))
            for column in range(4, len(_METRICS_HEADERS) + 1):
                ws.cell(row=row, column=column, value=NOT_MEASURED)
            row += 1
            continue

        other_agreements = [
            agreement_between(results[group_name], results[other])
            for other in names
            if other != group_name and results[other].status == "measured"
        ]
        other_agreements = [value for value in other_agreements if value is not None]
        mean_agreement = sum(other_agreements) / len(other_agreements) if other_agreements else None

        ws.cell(row=row, column=3, value="измерено")
        ws.cell(row=row, column=4, value=round(metrics.median_latency, 2)
                 if metrics.median_latency is not None else None)
        ws.cell(row=row, column=5, value=round(metrics.mean_latency, 2)
                 if metrics.mean_latency is not None else None)
        ws.cell(row=row, column=6, value=metrics.total_prompt_tokens)
        ws.cell(row=row, column=7, value=metrics.total_completion_tokens)
        ws.cell(row=row, column=8, value=_pct(metrics.valid_ratio))
        ws.cell(row=row, column=9, value=_pct(metrics.within_pool_ratio))
        ws.cell(row=row, column=10, value=_pct(metrics.reproducibility))
        ws.cell(row=row, column=11, value=_pct(mean_agreement))
        row += 1
    autosize_columns(ws)


def _write_profiles_sheet(
    wb: Workbook, results: dict[str, GroupResult], sample: list[str], repeats: int,
) -> None:
    ws = wb.create_sheet(SHEET_PROFILES)
    names = sorted(results)
    write_header_row(ws, ["Пользователь", "Повтор", *names])

    row = 2
    for user_id in sample:
        for repeat in range(repeats):
            ws.cell(row=row, column=1, value=user_id)
            ws.cell(row=row, column=2, value=repeat + 1)
            for column, group_name in enumerate(names, start=3):
                result = results[group_name]
                if result.status != "measured":
                    value = _not_measured_cell(summarize(result))
                else:
                    match = next(
                        (r for r in result.runs if r.user_id == user_id and r.repeat == repeat),
                        None,
                    )
                    if match is None:
                        value = "нет данных"
                    elif not match.order:
                        value = "трек пуст"
                    else:
                        value = " → ".join(match.order)
                ws.cell(row=row, column=column, value=value)
            row += 1
    autosize_columns(ws)


def _write_methodology_sheet(
    wb: Workbook,
    config: AppConfig,
    metrics_by_group: dict[str, GroupMetrics],
    sample: list[str],
    generated_at: datetime,
) -> None:
    ws = wb.create_sheet(SHEET_METHODOLOGY)
    write_header_row(ws, ["Параметр", "Значение"])

    rows = [
        ("Размер среза", len(sample)),
        ("Seed отбора", config.benchmark_sample_seed),
        ("Число повторов", config.benchmark_repeats),
        ("Дата прогона", generated_at.strftime("%Y-%m-%d %H:%M")),
    ]
    row = 2
    for label, value in rows:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Модели по группам").font = HEADER_FONT
    row += 1
    for group_name in sorted(metrics_by_group):
        metrics = metrics_by_group[group_name]
        ws.cell(row=row, column=1, value=group_name)
        ws.cell(
            row=row, column=2,
            value=(metrics.model or "—") if metrics.status == "measured" else _not_measured_cell(metrics),
        )
        row += 1
    autosize_columns(ws)


def write_benchmark_workbook(
    results: dict[str, GroupResult],
    config: AppConfig,
    sample: list[str],
    out_path: str | Path,
    generated_at: datetime | None = None,
) -> None:
    """Write the four-sheet comparison workbook: Сводка, Метрики, По
    профилям, Методика.

    `results` must include every configured group, `not_measured` ones
    included -- that is what turns an unreachable group into a labelled row
    on every sheet instead of a silently missing one.
    """
    metrics_by_group = {name: summarize(result) for name, result in results.items()}
    generated_at = generated_at or datetime.now()

    wb = Workbook()
    default_sheet = wb.active
    _write_summary_sheet(wb, metrics_by_group)
    _write_metrics_sheet(wb, metrics_by_group, results)
    _write_profiles_sheet(wb, results, sample, config.benchmark_repeats)
    _write_methodology_sheet(wb, config, metrics_by_group, sample, generated_at)
    wb.remove(default_sheet)
    wb.save(out_path)
